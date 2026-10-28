#!/usr/bin/env python3
"""Valida con lectores independientes las exportaciones que genera la prueba api_e2e.

    pip install openpyxl ezdxf pypdf ifcopenshell
    cargo test -p vigix-core --test api_e2e
    python tests/referencia/validar_exportaciones.py target/tmp/samples

Comprueba que cada archivo abre sin errores en un lector de terceros y que conserva lo que el
contrato exige: estado del documento en cada página, unidades en encabezados, huellas, perfil.
"""
import glob, json, os, sys, zipfile, hashlib

d = sys.argv[1] if len(sys.argv) > 1 else "target/tmp/samples"
fails = []


def check(cond, msg):
    print(("  OK  " if cond else "  FALLA ") + msg)
    if not cond:
        fails.append(msg)


def one(pattern):
    r = sorted(glob.glob(os.path.join(d, pattern)))
    assert r, f"no hay archivos {pattern} en {d}"
    return r[0]


# ---------------------------------------------------------------- PDF
import pypdf
for lang, banner, wm in (("es", "BORRADOR / NO APTO PARA CONSTRUCCI", "FREEMIUM"), ("en", "DRAFT / NOT FOR CONSTRUCTION", "FREEMIUM")):
    f = one(f"{lang}_*memoria*.pdf")
    r = pypdf.PdfReader(f)
    print(f"PDF {lang}: {os.path.basename(f)} ({len(r.pages)} páginas)")
    texts = [p.extract_text() for p in r.pages]
    check(len(r.pages) >= 5, "tiene varias páginas")
    check(all(banner in t for t in texts), "el banner de estado está en TODAS las páginas")
    check(all("%d" % (i + 1) in t and ("Página" in t or "Page" in t) for i, t in enumerate(texts)), "numeración de página en el pie")
    check(all(wm in t for t in texts), "marca de evaluación Freemium en todas las páginas")
    full = "\n".join(texts)
    for needle in (("DEMO-NO-NORMATIVO", "perfil de demostración citado"), ("SHA-256", "huella de datos"), ("bar", "unidades en cabeceras")):
        check(needle[0] in full, needle[1])
    check(any(p.mediabox.width > p.mediabox.height for p in r.pages), "hay páginas apaisadas (planos)")
    check("EJEMPLO-DEMO-001" in full, "identificador de proyecto en el documento")

# ---------------------------------------------------------------- XLSX
import openpyxl
f = one("es_*datos*.xlsx")
wb = openpyxl.load_workbook(f)
print("XLSX:", os.path.basename(f), wb.sheetnames)
check(len(wb.sheetnames) >= 10, "una hoja por tabla")
meta = {r[0]: r[1] for r in wb[wb.sheetnames[0]].iter_rows(min_row=2, values_only=True)}
check(any("BORRADOR" in str(v) for v in meta.values()), "la hoja de metadatos declara el estado BORRADOR")
check(any(len(str(v)) == 64 for v in meta.values()), "metadatos con huella SHA-256 de datos")
hdr = [c.value for c in wb["Hidráulica - tramos"][1]]
check(any("[bar]" in str(h) for h in hdr) and any("[L/min]" in str(h) for h in hdr) and any("[mm]" in str(h) for h in hdr), f"unidades en los encabezados hidráulicos: {hdr}")
nrows = wb["Hidráulica - tramos"].max_row
check(nrows > 5, f"tramos hidráulicos exportados ({nrows - 1})")
ws = wb["Incidencias"]
check(ws.max_row > 5, "incidencias exportadas")
check(all(c.data_type != "f" for row in ws.iter_rows() for c in row), "ninguna celda es una fórmula (sin inyección)")
check(wb["Parámetros normativos"].max_row > 5, "parámetros normativos con referencia")

# ---------------------------------------------------------------- CSV (ZIP)
f = one("es_*csv*.zip")
z = zipfile.ZipFile(f)
names = z.namelist()
print("CSV ZIP:", names)
check("MANIFEST.json" in names and "LEEME.txt" in names, "paquete con MANIFEST y LEEME")
man = json.loads(z.read("MANIFEST.json"))
check(all(hashlib.sha256(z.read(e["name"])).hexdigest() == e["sha256"] for e in man["files"]), "huellas SHA-256 del manifiesto coinciden")
check(z.read("issues.csv").startswith(b"\xef\xbb\xbf"), "CSV con BOM UTF-8")
import csv, io
rows = list(csv.reader(io.StringIO(z.read("hyd_links.csv").decode("utf-8-sig"))))
check("[bar]" in ",".join(rows[0]), "CSV con unidades en el encabezado")

# ---------------------------------------------------------------- DXF
import ezdxf
f = one("es_*plano_L1*.dxf")
doc = ezdxf.readfile(f)
aud = doc.audit()
print("DXF:", os.path.basename(f), doc.dxfversion)
check(len(aud.errors) == 0, f"ezdxf: auditoría sin errores ({len(aud.errors)})")
layers = {l.dxf.name for l in doc.layers}
check({"VIGIX-ARQ-MUROS", "VIGIX-ARQ-ZONAS", "VIGIX-EXT-PORTATIL", "VIGIX-AGUA-TUBERIA", "VIGIX-AGUA-DISPOSITIVO"} <= layers, "capas VIGIX-* presentes")
msp = doc.modelspace()
check(len(list(msp.query("INSERT"))) >= 8, "símbolos genéricos insertados como bloques")
check(any("BORRADOR" in e.dxf.text.replace("\\U+00D3", "O") or "BORRADOR" in e.dxf.text for e in msp.query("TEXT")), "texto de estado en el plano")

# ---------------------------------------------------------------- IFC
import ifcopenshell, ifcopenshell.geom, ifcopenshell.util.element as ue
f = one("es_*bim*.ifc")
m = ifcopenshell.open(f)
print("IFC:", os.path.basename(f), m.schema)
check(m.schema == "IFC4", "esquema IFC4")
check(len(m.by_type("IfcFireSuppressionTerminal")) >= 8 and len(m.by_type("IfcPipeSegment")) >= 6 and len(m.by_type("IfcSpace")) == 3, "terminales, tuberías y espacios")
settings = ifcopenshell.geom.settings()
bad = 0
for e in m.by_type("IfcProduct"):
    if e.Representation:
        try:
            ifcopenshell.geom.create_shape(settings, e)
        except Exception as ex:
            bad += 1
check(bad == 0, f"toda la geometría es válida ({bad} errores)")
t = m.by_type("IfcFireSuppressionTerminal")[0]
ps = ue.get_psets(t).get("Pset_VIGIX", {})
check("BORRADOR" in str(ps.get("DocumentStatus")) and ps.get("ProfileVerified") is False and len(str(ps.get("DataFingerprintSHA256"))) == 64, f"propiedades de trazabilidad: {ps}")

# ---------------------------------------------------------------- JSON
f = one("es_*datos*.json")
j = json.load(open(f, encoding="utf-8"))
check(j["format"] == "vigix.export/1" and j["meta"]["document_status"] == "draft" and j["meta"]["profile"]["demo"] is True, "JSON con formato, estado y perfil")
check(j["profile_package"]["profile_code"] == "DEMO-NO-NORMATIVO", "JSON incluye el paquete del perfil (reproducible)")

print()
if fails:
    print(f"{len(fails)} comprobaciones FALLARON")
    sys.exit(1)
print("Todas las exportaciones se validaron con lectores independientes.")
